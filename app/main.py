from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
import httpx
from fastapi import FastAPI, HTTPException

import os
from datetime import datetime
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook

from typing import Optional

app = FastAPI(title="Coastal Guard AI")

# Allow React (running on port 5173) to talk to this backend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/api/health")
def health():
    return {"status": "Coastal Guard AI backend is running"}

@app.get("/api/alerts")
def get_alerts():
    # Dummy data for now; real AI results come later
    return [
        {"id": 1, "location": "Marina Drive Ernakulam", "risk": "High", "type": "Rough waves"},
        {"id": 2, "location": "Kozhikode ", "risk": "Medium", "type": "Oil spill suspected"},
        {'id': 3, "location": "Kollam", "risk": "Low", "type": "Minor debris"},
        {"id": 4, "location": "Trivandrum", "risk": "High", "type": "Strong currents"},
    ]
    



KERALA_PLACES = {
    "Kochi":              {"lat": 9.9312,  "lon": 76.2673, "sea_lat": 9.95,  "sea_lon": 76.10},
    "Thiruvananthapuram": {"lat": 8.5241,  "lon": 76.9366, "sea_lat": 8.45,  "sea_lon": 76.85},
    "Kozhikode":          {"lat": 11.2588, "lon": 75.7804, "sea_lat": 11.26, "sea_lon": 75.65},
    "Alappuzha":          {"lat": 9.4981,  "lon": 76.3388, "sea_lat": 9.50,  "sea_lon": 76.20},
    "Kollam":             {"lat": 8.8932,  "lon": 76.6141, "sea_lat": 8.88,  "sea_lon": 76.50},
    "Kannur":             {"lat": 11.8745, "lon": 75.3704, "sea_lat": 11.88, "sea_lon": 75.25},
}

WEATHER_CODES = {
    0: "Clear sky", 1: "Mostly clear", 2: "Partly cloudy", 3: "Overcast",
    45: "Fog", 48: "Fog",
    51: "Light drizzle", 53: "Drizzle", 55: "Heavy drizzle",
    61: "Light rain", 63: "Rain", 65: "Heavy rain",
    80: "Rain showers", 81: "Rain showers", 82: "Violent rain showers",
    95: "Thunderstorm", 96: "Thunderstorm with hail", 99: "Thunderstorm with hail",
}

@app.get("/api/weather")
def get_weather(
    city: Optional[str] = None,
    lat: Optional[float] = None,
    lon: Optional[float] = None
):
    # 1. Determine latitude & longitude
    if lat is not None and lon is not None:
        target_lat = lat
        target_lon = lon
        sea_lat = lat
        sea_lon = lon - 0.15  # Shift slightly west into the sea for marine wave data
        display_name = "Current Location"
    else:
        # Default to Kochi if city not provided or not found
        city_name = city if (city and city in KERALA_PLACES) else "Kochi"
        place = KERALA_PLACES[city_name]
        target_lat = place["lat"]
        target_lon = place["lon"]
        sea_lat = place["sea_lat"]
        sea_lon = place["sea_lon"]
        display_name = f"{city_name}, Kerala"
    try:
        # 2. Fetch live atmospheric weather from Open-Meteo
        weather = httpx.get(
            "https://api.open-meteo.com/v1/forecast",
            params={
                "latitude": target_lat,
                "longitude": target_lon,
                "current": "temperature_2m,relative_humidity_2m,weather_code,wind_speed_10m,visibility",
                "timezone": "Asia/Kolkata",
            },
            timeout=10,
        ).json()["current"]
        # 3. Fetch marine wave height
        marine = httpx.get(
            "https://marine-api.open-meteo.com/v1/marine",
            params={
                "latitude": sea_lat,
                "longitude": sea_lon,
                "current": "wave_height",
                "timezone": "Asia/Kolkata",
            },
            timeout=10,
        ).json()
        wave = marine.get("current", {}).get("wave_height")
    except Exception:
        raise HTTPException(status_code=502, detail="Weather service unavailable")
    return {
        "location": display_name,
        "temp": round(weather["temperature_2m"]),
        "condition": WEATHER_CODES.get(weather["weather_code"], "Unknown"),
        "wave_height": wave if wave is not None else 0,
        "wind_speed": round(weather["wind_speed_10m"]),
        "humidity": weather["relative_humidity_2m"],
        "visibility": round((weather.get("visibility") or 0) / 1000, 1),
    }
    
    
    
# --- User Models ---
class RegisterRequest(BaseModel):
    name: str
    email: str
    phone: str
    role: str = "User" 
    location: str 
    password: str

class LoginRequest(BaseModel):
    email: str
    password: str

EXCEL_FILE = "users.xlsx"

# Helper: Append user (including password) to Excel
def append_user_to_excel(name: str, email: str, phone: str, role: str, location: str, password: str):
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Users"
        # Column headers including Password
        ws.append(["Full Name", "Email", "Phone", "Role", "Location", "Password", "Registered At"])
    else:
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws.append([name, email, phone, role, location, password, timestamp])
    wb.save(EXCEL_FILE)

# Helper: Search Excel to verify login credentials
def verify_user_in_excel(email: str, password: str):
    if not os.path.exists(EXCEL_FILE):
        return None, "No users registered yet"

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    # Check each row starting from row 2 (skipping header)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 6:
            continue

        user_name = row[0]
        user_email = str(row[1]).strip().lower() if row[1] else ""
        user_role = row[3] or "User"
        user_loc = row[4] or ""
        user_pwd = str(row[5]) if row[5] is not None else ""

        if user_email == email.strip().lower():
            if user_pwd == password:
                return {
                    "name": user_name,
                    "email": user_email,
                    "role": user_role,
                    "location": user_loc
                }, None
            else:
                return None, "Incorrect password"

    return None, "User not found with this email"



# 3. Register Endpoint
@app.post("/api/register")
def register(user: RegisterRequest):
    try:
        # Check if email is already taken
        if os.path.exists(EXCEL_FILE):
            wb = load_workbook(EXCEL_FILE)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) > 1 and str(row[1]).strip().lower() == user.email.strip().lower():
                    raise HTTPException(status_code=400, detail="Email is already registered")
        append_user_to_excel(
            user.name,
            user.email,
            user.phone,
            user.role,
            user.location,
            user.password
        )
        return {"status": "success", "message": "Registered successfully"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# --- 1. ADMIN USER REGISTRY ENDPOINT ---
@app.get("/api/admin/users")
def get_all_users():
    if not os.path.exists(EXCEL_FILE):
        return []
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    user_list = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not row or len(row) < 5:
            continue
        user_list.append({
            "id": idx,
            "name": row[0] or "N/A",
            "email": row[1] or "N/A",
            "phone": row[2] or "N/A",
            "role": row[3] or "User",
            "location": row[4] or "Kerala",
            "registered_at": row[6] if len(row) > 6 else "Recent"
        })
    return user_list


# --- 2. LOGIN ENDPOINT (WITH ADMIN BYPASS) ---
@app.post("/api/login")
def login(creds: LoginRequest):
    clean_email = creds.email.strip().lower()
    
    # MASTER ADMIN CHECK: username "admin" and password "admin123"
    if clean_email in ["admin", "admin@coastalguard.ai"] and creds.password == "admin123":
        return {
            "status": "success",
            "message": "Admin login successful",
            "user": {
                "name": "Super Admin",
                "email": "admin@coastalguard.ai",
                "role": "admin",
                "location": "Kerala State Command (HQ)"
            }
        }
    # Standard User / Volunteer check from Excel
    user_data, error = verify_user_in_excel(creds.email, creds.password)
    if error:
        raise HTTPException(status_code=401, detail=error)
    return {"status": "success", "message": "Login successful", "user": user_data}


